"""Document extraction service + upload API contract tests.

Covers:
    1. PDF extraction
    2. DOCX extraction
    3. TXT extraction
    4. unsupported extension / magic-byte rejection
    5. empty document rejection
    6. extracted text flowing through the existing analysis pipeline
    7. multipart route contract (extract + analyze), incl. size/character caps

Route-level tests invoke the FastAPI route functions directly with a real
starlette ``UploadFile``. This keeps the shared tests' temp SQLite file away
from TestClient's background portal (which otherwise holds the file open on
Windows and breaks conftest's per-test reset).
"""

from __future__ import annotations

import io

import pytest
from fastapi import HTTPException
from starlette.datastructures import Headers, UploadFile

from app.api.routes import documents
from app.config import get_settings
from app.services.document_extractor.extractor import (
    EmptyDocumentError,
    UnsupportedDocumentTypeError,
    extract_text,
)
from app.services.normalization.ontology import get_ontology
from app.services.pipeline import AnalysisPipeline

NARRATIVE = (
    "During routine flange maintenance on the gas line, the fitter opened the "
    "flange before verifying zero energy, resulting in a small gas release."
)


def _flatten(text: str) -> str:
    """Collapse arbitrary whitespace so line-wrapped PDF text compares cleanly."""
    return " ".join(text.split())


def _pdf_bytes() -> bytes:
    import fitz

    document = fitz.open()
    try:
        page = document.new_page()
        rect = fitz.Rect(50, 50, page.rect.width - 50, page.rect.height - 50)
        page.insert_textbox(rect, NARRATIVE, fontsize=10)
        return document.tobytes()
    finally:
        document.close()


def _docx_bytes() -> bytes:
    from docx import Document

    buf = io.BytesIO()
    document = Document()
    document.add_paragraph(NARRATIVE)
    document.add_paragraph("No injury occurred.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Location: flange area"
    document.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------- extraction


def test_txt_extraction():
    doc = extract_text("report.txt", NARRATIVE.encode("utf-8"))
    assert doc.file_type == "txt"
    assert NARRATIVE in doc.text
    assert doc.character_count == len(NARRATIVE)


def test_pdf_extraction():
    doc = extract_text("report.pdf", _pdf_bytes())
    assert doc.file_type == "pdf"
    assert NARRATIVE in _flatten(doc.text)
    assert doc.character_count > 0


def test_docx_extraction_preserves_paragraphs_and_tables():
    doc = extract_text("report.docx", _docx_bytes())
    assert doc.file_type == "docx"
    assert NARRATIVE in doc.text
    assert "No injury occurred." in doc.text
    assert "Location: flange area" in doc.text


# ----------------------------------------------------------------- reject


def test_unsupported_extension_rejected():
    with pytest.raises(UnsupportedDocumentTypeError):
        extract_text("report.png", b"%PDF-whatever")


def test_magic_bytes_mismatch_rejected():
    # Declared PDF but not a PDF, declared DOCX but not a ZIP container.
    with pytest.raises(UnsupportedDocumentTypeError):
        extract_text("fake.pdf", b"just some text, not a pdf")
    with pytest.raises(UnsupportedDocumentTypeError):
        extract_text("fake.docx", b"plain text pretending to be a docx")


def test_empty_document_rejected():
    with pytest.raises(EmptyDocumentError):
        extract_text("empty.txt", b"")

    from docx import Document

    buf = io.BytesIO()
    document = Document()
    document.add_paragraph("")
    document.save(buf)
    with pytest.raises(EmptyDocumentError):
        extract_text("empty.docx", buf.getvalue())


def test_blank_pdf_with_no_text_rejected():
    import fitz

    document = fitz.open()
    try:
        document.new_page()
        data = document.tobytes()
    finally:
        document.close()
    with pytest.raises(EmptyDocumentError):
        extract_text("blank.pdf", data)


# ----------------------------------------------------- pipeline integration


def test_extracted_text_reaches_existing_pipeline():
    doc = extract_text("report.txt", NARRATIVE.encode("utf-8"))
    pipeline = AnalysisPipeline(get_ontology(), get_settings())
    result = pipeline.analyze("OBS-TEST", doc.text, provider="rules")

    assert result.event.narrative == doc.text
    assert result.event.barrier == "energy_isolation"
    assert result.event.barrier_state == "not_verified"
    assert result.event.energy == "pressurized_gas"
    assert result.event.exposure == "uncontrolled_gas_release"
    assert result.event.sif != ""


# ------------------------------------------------------------- API contract


def _upload(filename: str, data: bytes, content_type: str = "application/octet-stream") -> UploadFile:
    return UploadFile(
        file=io.BytesIO(data),
        filename=filename,
        headers=Headers({"content-type": content_type}),
    )


def test_extract_endpoint_returns_text():
    body = documents.extract_document(
        _upload("report.txt", NARRATIVE.encode("utf-8"), "text/plain")
    )
    assert body["filename"] == "report.txt"
    assert body["file_type"] == "txt"
    assert body["text"] == NARRATIVE
    assert body["character_count"] == len(NARRATIVE)


def test_extract_pdf_endpoint():
    body = documents.extract_document(_upload("report.pdf", _pdf_bytes(), "application/pdf"))
    assert body["file_type"] == "pdf"
    assert NARRATIVE in _flatten(body["text"])


def test_extract_unsupported_extension_415():
    with pytest.raises(HTTPException) as exc:
        documents.extract_document(_upload("photo.png", b"not really png", "image/png"))
    assert exc.value.status_code == 415


def test_extract_empty_document_422():
    with pytest.raises(HTTPException) as exc:
        documents.extract_document(_upload("empty.txt", b"   \n \n", "text/plain"))
    assert exc.value.status_code == 422


def test_extract_character_budget_422():
    long_text = ("flange maintenance described here. " * 400)[:4000]
    with pytest.raises(HTTPException) as exc:
        documents.extract_document(_upload("long.txt", long_text.encode("utf-8"), "text/plain"))
    assert exc.value.status_code == 422
    assert "exceeding" in exc.value.detail


def test_extract_size_limit_413(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "max_upload_bytes", 8)
    with pytest.raises(HTTPException) as exc:
        documents.extract_document(_upload("big.txt", b"x" * 1024, "text/plain"))
    assert exc.value.status_code == 413


def test_analyze_document_endpoint_runs_pipeline():
    res = documents.analyze_document(_upload("report.docx", _docx_bytes()))
    assert res.id
    assert res.report_id.startswith("DOC-DOCX")
    assert res.event["barrier"] == "energy_isolation"
    assert res.event["barrier_state"] == "not_verified"
    assert res.event["narrative"] == extract_text("report.docx", _docx_bytes()).text


def test_analyze_document_too_little_text_422():
    with pytest.raises(HTTPException) as exc:
        documents.analyze_document(_upload("tiny.txt", b"no", "text/plain"))
    assert exc.value.status_code == 422