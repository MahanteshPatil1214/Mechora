"""Document text extraction for uploaded HSE reports.

Supported formats:
    - PDF   (PyMuPDF / ``fitz``)
    - DOCX  (python-docx)
    - TXT   (utf-8 text, tolerant decode)

The extractor preserves the exact extracted wording as much as the source
format allows (PDF layout order may differ from visual reading order; DOCX
paragraph and table text is concatenated line-by-line). No OCR, no structure
inference, no summarization.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt"})


class UnsupportedDocumentTypeError(ValueError):
    """Raised when a file extension (or magic bytes) is not supported."""


class EmptyDocumentError(ValueError):
    """Raised when a supported document yields no extractable text."""


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    file_type: str
    text: str

    @property
    def character_count(self) -> int:
        return len(self.text)


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    """Extract normalized plain text from an uploaded document.

    Raises:
        UnsupportedDocumentTypeError: unknown / mismatched file type.
        EmptyDocumentError: the document contained no extractable text.
    """
    if not data:
        raise EmptyDocumentError("uploaded file is empty")

    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentTypeError(
            f"unsupported file type '{suffix or 'none'}'; "
            f"accepted: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    _assert_magic_bytes(suffix, data)

    if suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".docx":
        text = _extract_docx(data)
    else:
        text = _extract_txt(data)

    text = _normalize(text)
    if not text.strip():
        raise EmptyDocumentError("document contains no extractable text")
    return ExtractedDocument(filename=filename, file_type=suffix.lstrip("."), text=text)


def _assert_magic_bytes(suffix: str, data: bytes) -> None:
    """Reject a file whose declared extension does not match its content."""
    is_pdf = data[:5] == b"%PDF-"
    is_zip = data[:4] in (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")

    if suffix == ".pdf" and not is_pdf:
        raise UnsupportedDocumentTypeError(
            "file does not look like a PDF (missing %PDF header)"
        )
    if suffix == ".docx" and not is_zip:
        raise UnsupportedDocumentTypeError(
            "file does not look like a DOCX (not a ZIP/OOXML container)"
        )


def _extract_pdf(data: bytes) -> str:
    import fitz  # PyMuPDF

    with fitz.open(stream=data, filetype="pdf") as document:
        pages = [(page.get_text("text") or "") for page in document]
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _normalize(text: str) -> str:
    """Collapse runaway blank-line runs while preserving wording."""
    lines: list[str] = []
    blank_run = 0
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            blank_run += 1
            if blank_run == 1:
                lines.append(line)
        else:
            blank_run = 0
            lines.append(line)
    return "\n".join(lines).strip()